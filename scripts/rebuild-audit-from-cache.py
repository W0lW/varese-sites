#!/usr/bin/env python3
"""Rebuild the Varese Sites workbook from the completed audit artifact, offline.
Search-result classifications are deliberately conservative: directory listings are
not treated as official websites, and missing search cache is never treated as proof
that a business has no website.
"""
import csv, json, re, unicodedata
from pathlib import Path
from urllib.parse import urlparse
from collections import Counter
from datetime import datetime, timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path.cwd()
REPORT = ROOT / "audit-report"
CACHE_PATH = REPORT / "web-search-cache.json"
PAGES_PATH = REPORT / "all-pages.csv"
MANIFEST_PATH = ROOT / "varese_sites_missing_manifest.csv"
CORRECTIONS_PATH = ROOT / "audit-template-corrections.csv"
OUT = REPORT / "audit-completo-attivita.xlsx"

def read_csv(path):
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))

def clean(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()

def norm(value):
    value = unicodedata.normalize("NFKD", str(value or "").lower())
    value = "".join(c for c in value if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", "", value)

DIRECTORY_HOST_MARKERS = (
    "facebook.com", "instagram.com", "linkedin.com", "tiktok.com", "youtube.com",
    "twitter.com", "x.com", "threads.net", "tripadvisor.", "paginegialle.",
    "paginebianche.", "firmania.", "restaurantguru.", "reteimprese.", "pos.do",
    "italiarecensioni.", "misterimprese.", "yelp.", "oraridiapertura24.",
    "informazione-aziende.", "cylex-", "aziende.virgilio.", "virgilio.it",
    "fresha.com", "booking.com", "findglocal.com", "thepubcrawlapp.com",
    "near-place.com", "ameve.it", "ufficiocamerale.", "beautynailhairsalons.com",
    "worldplaces.me", "mapcarta.com", "dnb.com", "cybo.com", "italy724.info",
    "coobiz.it", "localitybiz.it", "openalacarte.com", "sluurpy.it",
    "bed-and-breakfast.it", "mapy.com", "infobel.", "menuweb.menu",
    "empresite.it", "altraopinione.org", "eatoutmap.com", "negozi.sisal.com",
    "trip.com", "myadj.it", "registroaziende.it", "treatwell.", "justeat.",
    "deliveroo.", "glovoapp.", "ubereats.", "thefork.", "quandoo.", "kompass.",
    "tuttocitta.it", "impresaitalia.info", "prontoimprese.", "google.com",
    "maps.google.", "goo.gl", "wikipedia.org", "foursquare.com", "mapquest.com",
    "yellowpages.", "yelp.", "it.indeed.com", "glassdoor.", "crunchbase.com",
    "opendi.", "pagine.it", "orari-di-apertura.", "trovaprezzi.", "doveconviene."
)
EMAIL_RE = re.compile(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", re.I)
PHONE_RE = re.compile(r"(?<!\d)(?:(?:\+39|0039)[\s./-]*)?(?:0\d{1,3}|3\d{2})[\s./-]?\d{3,4}[\s./-]?\d{3,4}(?!\d)")

STOP_TOKENS = {
    "centro", "studi", "studio", "bar", "caffe", "cafe", "ristorante", "ristoro",
    "hotel", "albergo", "pizzeria", "ristoranti", "negozio", "negozi", "servizi",
    "service", "services", "societa", "srl", "snc", "sas", "spa", "di", "da",
    "del", "della", "delle", "dei", "degli", "per", "the", "and", "e", "la",
    "il", "lo", "le", "un", "una", "azienda", "attivita", "varese", "busto",
    "arsizio", "gallarate", "luino", "tradate", "italia", "italy"
}

def is_directory(url):
    host = urlparse(url).netloc.lower().removeprefix("www.")
    return any(marker in host for marker in DIRECTORY_HOST_MARKERS)

def extract_contacts(sources):
    phones, emails, evidence = [], [], []
    for source in sources:
        desc = clean(source.get("description"))
        url = clean(source.get("url"))
        found_here = False
        for email in EMAIL_RE.findall(desc):
            email = email.lower()
            if email not in emails:
                emails.append(email)
                found_here = True
        for raw in PHONE_RE.findall(desc):
            digits = re.sub(r"\D", "", raw)
            if digits.startswith("0039"):
                digits = digits[4:]
            elif digits.startswith("39") and len(digits) > 10:
                digits = digits[2:]
            if 9 <= len(digits) <= 10 and digits not in phones:
                phones.append(digits)
                found_here = True
        if found_here and url and url not in evidence:
            evidence.append(url)
    return phones[:4], emails[:4], evidence[:6]

def classify_presence(name, sources, cached):
    if not cached:
        return ("DA VERIFICARE — ricerca non disponibile nella cache", "",
                "Il risultato della ricerca non è disponibile nel file cache: non equivale a prova di assenza del sito.")
    if not sources:
        return ("DA VERIFICARE — ricerca senza risultati", "",
                "Nessun risultato utile recuperato; ripetere la ricerca e controllare Maps/social.")
    tokens = [norm(t) for t in re.findall(r"[A-Za-zÀ-ÿ0-9]+", name)]
    tokens = [t for t in tokens if (len(t) >= 4 or (t.isdigit() and len(t) >= 2)) and t not in STOP_TOKENS]
    branded, other = [], []
    for source in sources:
        url = clean(source.get("url"))
        host = urlparse(url).netloc.lower().removeprefix("www.")
        if not host or is_directory(url):
            continue
        first_label = host.split(".")[0]
        host_norm = norm(first_label)
        if any(token in host_norm for token in tokens):
            branded.append(source)
        else:
            other.append(source)
    if branded:
        return ("SITO PROBABILE — ESCLUDERE SOLO DOPO CONFERMA", branded[0].get("url", ""),
                "Dominio non classificato come directory che richiama una parte distintiva del nome. Aprire il sito e verificare che appartenga davvero all'attività.")
    if other:
        return ("INCERTO — CONTROLLO MANUALE NECESSARIO", other[0].get("url", ""),
                "È emerso un dominio non chiaramente riconducibile all'attività; verificare manualmente prima di includere/escludere la prospect.")
    return ("NESSUN SITO UFFICIALE INDIVIDUATO — PROSPECT DA CONFERMARE", "",
            "Nei risultati recuperati compaiono solo directory, social o piattaforme terze. Controllare Maps, social e ricerca manuale prima del contatto.")

pages = read_csv(PAGES_PATH)
manifest_rows = read_csv(MANIFEST_PATH)
corrections = read_csv(CORRECTIONS_PATH)
with CACHE_PATH.open("r", encoding="utf-8") as f:
    cache = json.load(f)
manifest = {clean(r.get("slug")): r for r in manifest_rows}
correction_map = {clean(r.get("slug")): r for r in corrections}

def business_name(page, m):
    name = clean(m.get("nome"))
    if name:
        return name
    name = clean(page.get("h1")) or clean(page.get("title"))
    name = re.split(r"\s+[|—–-]\s+", name)[0]
    return name or clean(page.get("slug"))

records, source_rows = [], []
status_counts = Counter()
template_counts = Counter()
flag_counts = Counter()
for page in pages:
    slug = clean(page.get("slug"))
    m = manifest.get(slug, {})
    has_manifest = slug in manifest
    name = business_name(page, m)
    comune, settore = clean(m.get("comune")), clean(m.get("settore"))
    query_parts = [f'"{name}"', comune or "provincia di Varese"]
    if settore:
        query_parts.append(settore)
    query = " ".join(query_parts)
    cached = cache.get(query)
    sources = cached.get("results", []) if cached else []
    if has_manifest:
        status, candidate, note = classify_presence(name, sources, cached)
    else:
        status, candidate, note = ("FUORI MANIFEST — NON TRATTARE COME PROSPECT", "",
            "Questa pagina non corrisponde a una voce del manifest attività; non inserirla nella lista commerciale finché nome, comune e attività non siano verificati.")
    status_counts[status] += 1
    web_phones, web_emails, contact_evidence = extract_contacts(sources)
    rec = {
        "slug": slug, "name": name, "comune": comune, "settore": settore, "query": query, "has_manifest": has_manifest,
        "status": status, "candidate": candidate, "note": note, "sources": sources,
        "phone": clean(m.get("telefono")), "email": clean(m.get("email")),
        "web_phones": "; ".join(web_phones), "web_emails": "; ".join(web_emails),
        "contact_evidence": "; ".join(contact_evidence),
        "address": clean(m.get("indirizzo")), "social": clean(m.get("sito_social_trovato")),
        "demo": clean(m.get("url_demo")) or f"https://w0lw.github.io/varese-sites/sites/{slug}/",
        "page": page, "manifest": m, "correction": correction_map.get(slug, {}),
        "search_state": (cached or {}).get("status", "NON PRESENTE NELLA CACHE")
    }
    records.append(rec)
    template_counts[clean(page.get("template")) or "MISSING"] += 1
    for flag in clean(page.get("flags")).split(";"):
        if flag:
            flag_counts[flag] += 1
    if sources:
        for i, s in enumerate(sources, 1):
            source_rows.append([
                slug, name, comune, settore, query, s.get("engine", "DuckDuckGo"),
                cached.get("status", ""), i, clean(s.get("title")), clean(s.get("url")),
                clean(s.get("description")), status, candidate, note,
                "Aprire la fonte; verificare nome, settore, contatti e appartenenza del dominio."
            ])
    else:
        source_rows.append([
            slug, name, comune, settore, query, "DuckDuckGo", rec["search_state"], "", "", "", "",
            status, candidate, note,
            "Ripetere la ricerca: assenza di risultati in cache non dimostra assenza di sito."
        ])

wb = Workbook()
summary = wb.active
summary.title = "Riepilogo"
summary.append(["Voce", "Valore"])
summary_data = [
    ("Generato (UTC)", datetime.now(timezone.utc).isoformat(timespec="seconds")),
    ("Pagine sottoposte a controllo", len(records)),
    ("Attività presenti nel manifest", len(manifest)),
    ("Pagine fuori manifest — escluse dalla lista prospect", sum(1 for r in records if not r["has_manifest"])),
    ("Ricerche con risultati in cache", sum(1 for r in records if r["sources"])),
    ("Sito probabile (dominio specifico; confermare)", sum(1 for r in records if r["status"].startswith("SITO PROBABILE"))),
    ("Nessun sito ufficiale individuato — prospect da confermare", sum(1 for r in records if r["status"].startswith("NESSUN SITO UFFICIALE"))),
    ("Controllo manuale (dominio ambiguo)", sum(1 for r in records if r["status"].startswith("INCERTO"))),
    ("Ricerca non disponibile/senza risultati utili", sum(1 for r in records if r["status"].startswith("DA VERIFICARE"))),
    ("Fonti web raccolte", sum(len(r["sources"]) for r in records)),
    ("Ricerche riuscite presenti in cache", len(cache)),
    ("Metodo", "Riclassificazione conservativa dei risultati del precedente audit; nessuna nuova ricerca web eseguita."),
    ("Limite importante", "Non è una verifica definitiva: i candidati senza sito vanno controllati su Google/Maps/social; un dominio aziendale può non contenere il nome e finire nel controllo manuale."),
    ("Contatti", "Telefono, email, indirizzo e social provengono dal manifest iniziale e possono essere incompleti o da aggiornare."),
    ("Prospect con telefono", sum(1 for r in records if r["has_manifest"] and r["status"].startswith("NESSUN SITO UFFICIALE") and r["phone"])),
    ("Prospect con email", sum(1 for r in records if r["has_manifest"] and r["status"].startswith("NESSUN SITO UFFICIALE") and r["email"])),
    ("Prospect con telefono o email nel manifest", sum(1 for r in records if r["has_manifest"] and r["status"].startswith("NESSUN SITO UFFICIALE") and (r["phone"] or r["email"]))),
    ("Prospect con contatto trovato negli snippet web (da confermare)", sum(1 for r in records if r["has_manifest"] and r["status"].startswith("NESSUN SITO UFFICIALE") and (r["web_phones"] or r["web_emails"]))),
    ("Prospect con almeno un contatto nel manifest o negli snippet", sum(1 for r in records if r["has_manifest"] and r["status"].startswith("NESSUN SITO UFFICIALE") and (r["phone"] or r["email"] or r["web_phones"] or r["web_emails"]))),
    ("Esito template", "I flag HTML sono controlli automatici: richiedono verifica visiva e dei contenuti. La presenza di testo demo non prova da sola ogni singolo errore."),
]
for row in summary_data:
    summary.append(row)
summary.append([])
summary.append(["Anomalia strutturale", "Numero pagine"])
for k, v in sorted(flag_counts.items(), key=lambda x: (-x[1], x[0])):
    summary.append([k, v])
summary.append([])
summary.append(["Template assegnato", "Numero pagine"])
for k, v in sorted(template_counts.items(), key=lambda x: (-x[1], x[0])):
    summary.append([k, v])

activities = wb.create_sheet("Attività")
activities.append([
    "attività", "comune", "settore", "telefono manifest", "email manifest", "telefono da snippet web (da verificare)",
    "email da snippet web (da verificare)", "fonti dei contatti estratti", "indirizzo manifest",
    "social/sito trovato nel manifest", "stato sito da ricerca", "URL probabile sito ufficiale",
    "nota verifica sito", "query ricerca", "URL sito demo", "slug", "template assegnato",
    "titolo pagina", "H1", "anomalie automatiche", "decisione audit template", "motivazione audit template"
])
for r in records:
    p, c = r["page"], r["correction"]
    activities.append([
        r["name"], r["comune"], r["settore"], r["phone"], r["email"], r["web_phones"], r["web_emails"],
        r["contact_evidence"], r["address"], r["social"], r["status"], r["candidate"], r["note"], r["query"], r["demo"], r["slug"],
        clean(p.get("templateTitle")) or clean(p.get("template")), clean(p.get("title")),
        clean(p.get("h1")), clean(p.get("flags")), clean(c.get("decisione_audit")), clean(c.get("motivo_audit"))
    ])

research = wb.create_sheet("Ricerca web")
research.append([
    "slug", "attività", "comune", "settore", "query", "motore", "esito ricerca", "fonte #",
    "titolo fonte", "URL fonte", "estratto fonte", "classificazione sito", "URL probabile sito",
    "nota", "azione"
])
for row in source_rows:
    research.append(row)

prospects = wb.create_sheet("Prospect senza sito")
prospects.append([
    "attività", "comune", "settore", "telefono manifest", "email manifest", "telefono da snippet web (da verificare)",
    "email da snippet web (da verificare)", "fonti contatto web", "indirizzo", "social/sito noto",
    "stato", "fonte directory/social (non sito ufficiale)", "query", "sito demo", "prossima azione"
])
for r in sorted(records, key=lambda x: (x["comune"], x["name"])):
    if r["has_manifest"] and r["status"].startswith("NESSUN SITO UFFICIALE"):
        evidence = r["sources"][0].get("url", "") if r["sources"] else ""
        prospects.append([
            r["name"], r["comune"], r["settore"], r["phone"], r["email"], r["web_phones"],
            r["web_emails"], r["contact_evidence"], r["address"], r["social"], r["status"], evidence,
            r["query"], r["demo"],
            "Confermare contatti e assenza di sito ufficiale; solo dopo proporre il sito demo."
        ])

contact_sheet = wb.create_sheet("Contatti da verificare")
contact_sheet.append(["attività", "comune", "settore", "telefono manifest", "email manifest",
                      "telefono da snippet web", "email da snippet web", "fonti contatto web",
                      "stato sito", "sito demo"])
for r in sorted(records, key=lambda x: (x["comune"], x["name"])):
    if r["has_manifest"] and r["status"].startswith("NESSUN SITO UFFICIALE") and (r["phone"] or r["email"] or r["web_phones"] or r["web_emails"]):
        contact_sheet.append([r["name"], r["comune"], r["settore"], r["phone"], r["email"],
                              r["web_phones"], r["web_emails"], r["contact_evidence"], r["status"], r["demo"]])

exclude = wb.create_sheet("Sito probabile - verifica")
exclude.append([
    "attività", "comune", "settore", "telefono", "email", "indirizzo", "stato",
    "URL probabile sito ufficiale", "query", "sito demo", "nota"
])
for r in sorted(records, key=lambda x: (x["comune"], x["name"])):
    if r["has_manifest"] and r["status"].startswith("SITO PROBABILE"):
        exclude.append([r["name"], r["comune"], r["settore"], r["phone"], r["email"],
                        r["address"], r["status"], r["candidate"], r["query"], r["demo"], r["note"]])

manual = wb.create_sheet("Controllo manuale")
manual.append([
    "attività", "comune", "settore", "telefono", "email", "indirizzo", "social/sito noto",
    "stato", "URL da controllare", "query", "sito demo", "motivo/azione"
])
for r in sorted(records, key=lambda x: (x["comune"], x["name"])):
    if r["has_manifest"] and r["status"].startswith(("INCERTO", "DA VERIFICARE")):
        manual.append([r["name"], r["comune"], r["settore"], r["phone"], r["email"],
                       r["address"], r["social"], r["status"], r["candidate"], r["query"],
                       r["demo"], r["note"]])

outside = wb.create_sheet("Fuori manifest")
outside.append(["attività dalla pagina", "slug", "comune pagina", "settore pagina", "titolo pagina", "H1", "URL pagina demo", "motivo"])
for r in sorted(records, key=lambda x: (x["name"])):
    if not r["has_manifest"]:
        p = r["page"]
        outside.append([r["name"], r["slug"], r["comune"], r["settore"], clean(p.get("title")),
                        clean(p.get("h1")), r["demo"],
                        "Non presente nel manifest iniziale: non considerare un lead senza verifica manuale."])

corr = wb.create_sheet("Correzioni template")
if corrections:
    corr.append(list(corrections[0].keys()))
    for row in corrections:
        corr.append([row.get(k, "") for k in corrections[0].keys()])
else:
    corr.append(["Nessuna correzione template registrata"])

templates = wb.create_sheet("Distribuzione template")
templates.append(["Template ID", "Numero pagine", "Titolo/label"])
for key, count in sorted(template_counts.items(), key=lambda x: (-x[1], x[0])):
    label = next((clean(p.get("templateTitle")) for p in pages if clean(p.get("template")) == key and clean(p.get("templateTitle"))), "")
    templates.append([key, count, label])

for ws in wb.worksheets:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="20283A")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        max_len = 12
        for cell in ws[letter][:min(ws.max_row, 100)]:
            max_len = max(max_len, min(len(str(cell.value or "")), 42))
        ws.column_dimensions[letter].width = min(max_len + 2, 48)

for sheet, widths in {
    "Riepilogo": {"A": 62, "B": 100},
    "Attività": {"A": 32, "B": 22, "C": 42, "D": 22, "E": 34, "F": 34, "G": 42, "H": 52, "I": 60, "J": 72, "K": 64, "L": 54, "M": 30, "N": 30, "O": 44, "P": 34, "Q": 34, "R": 28, "S": 75},
    "Ricerca web": {"A": 30, "B": 32, "C": 22, "D": 40, "E": 64, "F": 20, "G": 36, "H": 10, "I": 48, "J": 60, "K": 75, "L": 52, "M": 60, "N": 75, "O": 72},
    "Prospect senza sito": {"A": 34, "B": 22, "C": 40, "D": 22, "E": 34, "F": 26, "G": 32, "H": 60, "I": 34, "J": 42, "K": 50, "L": 62, "M": 64, "N": 54, "O": 76},
    "Contatti da verificare": {"A": 34, "B": 22, "C": 40, "D": 22, "E": 34, "F": 30, "G": 34, "H": 80, "I": 52, "J": 54},
    "Sito probabile - verifica": {"A": 34, "B": 22, "C": 40, "D": 22, "E": 34, "F": 34, "G": 50, "H": 60, "I": 64, "J": 54, "K": 75},
    "Controllo manuale": {"A": 34, "B": 22, "C": 40, "D": 22, "E": 34, "F": 34, "G": 42, "H": 50, "I": 60, "J": 64, "K": 54, "L": 78},
}.items():
    for col, width in widths.items():
        wb[sheet].column_dimensions[col].width = width

wb.save(OUT)
print(json.dumps({
    "output": str(OUT), "activities": len(records), "cached_searches": len(cache),
    "sources": sum(len(r["sources"]) for r in records), "classification_counts": dict(status_counts),
    "prospects_to_confirm": sum(1 for r in records if r["status"].startswith("NESSUN SITO UFFICIALE")),
    "probable_official_domains": sum(1 for r in records if r["status"].startswith("SITO PROBABILE")),
    "manual_review": sum(1 for r in records if r["status"].startswith(("INCERTO", "DA VERIFICARE")))
}, ensure_ascii=False, indent=2))
