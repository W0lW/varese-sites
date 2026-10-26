#!/usr/bin/env python3
"""Create a traceable Excel audit and perform one real web search per business.

Brave Search is used if BRAVE_SEARCH_API_KEY exists; otherwise the DDGS library
queries DuckDuckGo. Search results are evidence leads, not automatic proof.
"""
import csv
import json
import os
import re
import time
import urllib.parse
import urllib.request
from urllib.parse import urlparse
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

try:
    from ddgs import DDGS
except ImportError:
    DDGS = None

ROOT = Path.cwd()
REPORT = ROOT / "audit-report"
REPORT.mkdir(exist_ok=True)
PAGES_CSV = REPORT / "all-pages.csv"
CORRECTIONS_CSV = ROOT / "audit-template-corrections.csv"
MANIFEST_CSV = ROOT / "varese_sites_missing_manifest.csv"
CACHE_PATH = REPORT / "web-search-cache.json"
OUT = REPORT / "audit-completo-attivita.xlsx"
API_KEY = os.getenv("BRAVE_SEARCH_API_KEY", "").strip()
MAX_SEARCHES = int(os.getenv("AUDIT_MAX_SEARCHES") or "2500")
PAUSE = float(os.getenv("AUDIT_SEARCH_PAUSE_SECONDS") or "0.5")
FLUSH_EVERY = 20

def read_csv(path):
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))

def clean(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()

def business_name(page, manifest):
    name = clean(manifest.get("nome"))
    if name:
        return name
    name = clean(page.get("h1")) or clean(page.get("title"))
    name = re.split(r"\s+[|—–-]\s+", name)[0]
    name = re.sub(r"^(home|benvenuti da|benvenuti a|welcome to)\s+", "", name, flags=re.I)
    return clean(name) or clean(page.get("slug"))

EXCLUDED_HOSTS = (
    "facebook.com", "instagram.com", "linkedin.com", "tiktok.com", "youtube.com",
    "x.com", "twitter.com", "tripadvisor.", "paginegialle.", "paginebianche.",
    "virgilio.it", "tuttocitta.it", "yelp.", "misterimprese.", "aziende.it",
    "treatwell.", "booking.com", "thefork.", "quandoo.", "google.com",
    "maps.google.", "goo.gl", "wikipedia.org", "kompass.", "cylex.", "infobel.",
    "foursquare.com", "justeat.", "deliveroo.", "glovoapp.", "ubereats.",
    "prontoimprese.", "impresaitalia.info", "registroimprese."
)

def classify_website_presence(sources, name, status):
    if not sources:
        return ("DA VERIFICARE MANUALMENTE — nessun risultato non prova l'assenza di sito", "", "Ricerca senza fonti utilizzabili; ripetere con altre query o controllare Maps/social.")
    name_tokens = [t.lower() for t in re.findall(r"[a-z0-9]+", name) if len(t) >= 3]
    directory_sources = []
    likely_sites = []
    other_sources = []
    for source in sources:
        url = clean(source.get("url"))
        title = clean(source.get("title"))
        desc = clean(source.get("description"))
        host = urlparse(url).netloc.lower().removeprefix("www.")
        text_blob = (title + " " + desc + " " + host).lower()
        if any(domain in host for domain in EXCLUDED_HOSTS):
            directory_sources.append(source)
            continue
        # Do not treat every unrelated result as the business's own site.
        token_hits = sum(1 for token in name_tokens if token in text_blob)
        if token_hits >= 1:
            likely_sites.append(source)
        else:
            other_sources.append(source)
    if likely_sites:
        return ("SITO WEB PROBABILE — ESCLUDERE DALLE PROSPECT FINO A CONTROLLO", likely_sites[0].get("url", ""),
                "Risultato su dominio non classificato come directory/social che richiama il nome dell'attività; confermare che sia il sito ufficiale e funzionante.")
    if directory_sources and not other_sources:
        return ("NESSUN SITO UFFICIALE INDIVIDUATO — POTENZIALE PROSPECT DA CONFERMARE", directory_sources[0].get("url", ""),
                "Sono emersi solo profili social, directory o piattaforme terze; verificare manualmente prima del contatto.")
    return ("INCERTO — CONTROLLO MANUALE NECESSARIO", (other_sources or directory_sources)[0].get("url", ""),
            "Risultati ambigui o non sufficienti per stabilire se esista un sito ufficiale.")

def write_cache(cache):
    CACHE_PATH.write_text(json.dumps(cache, ensure_ascii=False, indent=2), encoding="utf-8")

def web_search(query, cache):
    cached = cache.get(query)
    if cached and cached.get("status") == "RICERCA ESEGUITA" and cached.get("results"):
        return cached, False

    results = []
    status = "RICERCA ESEGUITA"
    if API_KEY:
        url = "https://api.search.brave.com/res/v1/web/search?" + urllib.parse.urlencode({
            "q": query, "count": 5, "country": "IT", "search_lang": "it", "safesearch": "moderate"
        })
        req = urllib.request.Request(url, headers={
            "Accept": "application/json",
            "X-Subscription-Token": API_KEY,
            "User-Agent": "VareseSitesAudit/1.0"
        })
        try:
            with urllib.request.urlopen(req, timeout=25) as response:
                payload = json.loads(response.read().decode("utf-8"))
            for item in payload.get("web", {}).get("results", [])[:5]:
                results.append({
                    "title": clean(item.get("title")),
                    "url": clean(item.get("url")),
                    "description": clean(item.get("description")),
                    "engine": "Brave Search",
                })
        except Exception as exc:
            status = "ERRORE RICERCA: " + clean(str(exc))[:220]
    elif DDGS is not None:
        last_error = ""
        for attempt in range(2):
            try:
                raw = DDGS().text(query, region="it-it", safesearch="moderate", max_results=5)
                for item in raw or []:
                    href = clean(item.get("href") or item.get("url"))
                    if href.startswith("http"):
                        results.append({
                            "title": clean(item.get("title")),
                            "url": href,
                            "description": clean(item.get("body") or item.get("description")),
                            "engine": "DuckDuckGo",
                        })
                break
            except Exception as exc:
                last_error = clean(str(exc))[:220]
                if attempt == 0:
                    time.sleep(2)
        if not results and last_error:
            status = "ERRORE RICERCA DuckDuckGo: " + last_error
        elif not results:
            status = "NESSUN RISULTATO"
    else:
        status = "NON ESEGUITA: nessun backend di ricerca disponibile"

    out = {"status": status, "results": results, "searched_at_utc": datetime.now(timezone.utc).isoformat(timespec="seconds")}
    if status == "RICERCA ESEGUITA" and results:
        cache[query] = out
    time.sleep(PAUSE)
    return out, True

pages = read_csv(PAGES_CSV)
if not pages:
    raise SystemExit("Manca audit-report/all-pages.csv: eseguire prima node scripts/audit-3d-business-pages.mjs")
manifest_rows = read_csv(MANIFEST_CSV)
manifest_map = {clean(row.get("slug")): row for row in manifest_rows}
corrections = read_csv(CORRECTIONS_CSV)
correction_map = {clean(row.get("slug")): row for row in corrections}

try:
    cache = json.loads(CACHE_PATH.read_text(encoding="utf-8")) if CACHE_PATH.exists() else {}
except Exception:
    cache = {}

research_rows = []
website_audit = {}
searched_this_run = 0
businesses_with_results = 0
businesses_with_no_results = 0
businesses_with_errors = 0
source_count = 0

for index, page in enumerate(pages):
    slug = clean(page.get("slug"))
    manifest = manifest_map.get(slug, {})
    name = business_name(page, manifest)
    comune = clean(manifest.get("comune"))
    settore = clean(manifest.get("settore"))
    location = comune or "provincia di Varese"
    parts = [f'"{name}"', location]
    if settore:
        parts.append(settore)
    query = " ".join(parts)

    if searched_this_run >= MAX_SEARCHES and not (
        cache.get(query, {}).get("status") == "RICERCA ESEGUITA" and cache.get(query, {}).get("results")
    ):
        result = {"status": f"NON ESEGUITA: limite configurato {MAX_SEARCHES}", "results": []}
    else:
        result, was_searched = web_search(query, cache)
        if was_searched:
            searched_this_run += 1

    sources = result.get("results", [])
    presence, candidate_url, presence_note = classify_website_presence(sources, name, result.get("status", ""))
    website_audit[slug] = {"status": presence, "url": candidate_url, "note": presence_note, "query": query,
                           "sources": sources, "name": name, "comune": comune, "settore": settore}
    if sources:
        businesses_with_results += 1
        source_count += len(sources)
        for number, source in enumerate(sources, 1):
            research_rows.append({
                "slug": slug, "attività": name, "comune": comune, "settore": settore,
                "query ricerca": query, "motore": source.get("engine", "Web search"),
                "esito ricerca": result.get("status", ""),
                "fonte #": number, "titolo fonte": source.get("title", ""),
                "URL fonte": source.get("url", ""), "estratto fonte": source.get("description", ""),
                "testo del sito da confrontare": clean(page.get("visibleText"))[:1500],
                "esito verifica contenuti": "FONTE TROVATA — confronto puntuale ancora da fare",
                "stato sito ufficiale": presence, "URL candidato sito ufficiale": candidate_url,
                "nota verifica sito": presence_note,
                "azione": "Confrontare nome, settore, servizi, indirizzo, contatti e affermazioni specifiche con la fonte."
            })
    else:
        if "ERRORE" in result.get("status", ""):
            businesses_with_errors += 1
        else:
            businesses_with_no_results += 1
        research_rows.append({
            "slug": slug, "attività": name, "comune": comune, "settore": settore,
            "query ricerca": query, "motore": "Brave Search" if API_KEY else "DuckDuckGo",
            "esito ricerca": result.get("status", ""),
            "fonte #": "", "titolo fonte": "", "URL fonte": "", "estratto fonte": "",
            "testo del sito da confrontare": clean(page.get("visibleText"))[:1500],
            "esito verifica contenuti": "NON VERIFICATO: nessuna fonte recuperata",
            "stato sito ufficiale": presence, "URL candidato sito ufficiale": candidate_url,
            "nota verifica sito": presence_note,
            "azione": "Ripetere la ricerca o trovare una fonte ufficiale affidabile; non trattare l'assenza di risultati come prova di falsità."
        })

    if (index + 1) % FLUSH_EVERY == 0:
        write_cache(cache)
        print(json.dumps({
            "progress": index + 1, "total": len(pages),
            "searched_this_run": searched_this_run,
            "businesses_with_results": businesses_with_results,
            "businesses_with_errors": businesses_with_errors
        }, ensure_ascii=False), flush=True)

write_cache(cache)

wb = Workbook()
summary = wb.active
summary.title = "Riepilogo"
summary.append(["Voce", "Valore"])
summary_data = [
    ("Generato (UTC)", datetime.now(timezone.utc).isoformat(timespec="seconds")),
    ("Pagine censite", len(pages)),
    ("Attività con risultati web", businesses_with_results),
    ("Sito ufficiale probabile — da escludere dalle prospect", sum(1 for v in website_audit.values() if v["status"].startswith("SITO WEB PROBABILE"))),
    ("Nessun sito ufficiale individuato — prospect da confermare", sum(1 for v in website_audit.values() if v["status"].startswith("NESSUN SITO UFFICIALE"))),
    ("Presenza sito incerta / controllo manuale", sum(1 for v in website_audit.values() if v["status"].startswith(("DA VERIFICARE", "INCERTO")))),
    ("Ricerche effettivamente avviate in questa esecuzione", searched_this_run),
    ("Attività senza risultati", businesses_with_no_results),
    ("Attività con errori di ricerca", businesses_with_errors),
    ("Fonti web raccolte", source_count),
    ("Motore usato", "Brave Search API" if API_KEY else "DuckDuckGo (DDGS)"),
    ("Chiave API richiesta", "Sì, solo per Brave; DuckDuckGo viene usato senza chiave" ),
    ("Correzioni template già registrate", len(corrections)),
    ("Nota metodologica", "I risultati di ricerca e gli snippet sono evidenze da controllare; non equivalgono a verifica automatica di tutte le affermazioni del sito."),
    ("Stato", "Ricerca web effettuata dove indicato; confronto semantico dei contenuti da completare prima di dichiarare un'attività verificata."),
]
for row in summary_data:
    summary.append(list(row))
summary.append([])
summary.append(["Anomalia strutturale", "Numero pagine"])
flag_counts = Counter()
template_counts = Counter()
for page in pages:
    template_counts[clean(page.get("template")) or "MISSING"] += 1
    for flag in clean(page.get("flags")).split(";"):
        if flag:
            flag_counts[flag] += 1
for key, count in sorted(flag_counts.items(), key=lambda item: (-item[1], item[0])):
    summary.append([key, count])
summary.append([])
summary.append(["Template assegnato", "Numero pagine"])
for key, count in sorted(template_counts.items(), key=lambda item: (-item[1], item[0])):
    summary.append([key, count])

activities = wb.create_sheet("Attività")
activities.append([
    "slug", "attività", "comune", "settore da manifest", "titolo pagina", "H1",
    "template 3D", "URL sito", "descrizione SEO", "anomalie automatiche",
    "testo visibile estratto", "presenza sito web", "URL candidato sito ufficiale", "nota verifica sito",
    "template finale audit", "decisione audit template", "motivazione audit"
])
for page in pages:
    slug = clean(page.get("slug"))
    manifest = manifest_map.get(slug, {})
    correction = correction_map.get(slug, {})
    activities.append([
        slug, business_name(page, manifest), clean(manifest.get("comune")),
        clean(manifest.get("settore")), page.get("title", ""), page.get("h1", ""),
        page.get("templateTitle") or page.get("template", ""),
        "https://w0lw.github.io/varese-sites/sites/" + slug + "/",
        page.get("metaDescription", ""), page.get("flags", ""), page.get("visibleText", ""),
        website_audit.get(slug, {}).get("status", "DA VERIFICARE MANUALMENTE"),
        website_audit.get(slug, {}).get("url", ""), website_audit.get(slug, {}).get("note", ""),
        correction.get("template_finale_audit", ""), correction.get("decisione_audit", ""),
        correction.get("motivo_audit", "")
    ])

research = wb.create_sheet("Ricerca web")
research_headers = [
    "slug", "attività", "comune", "settore", "query ricerca", "motore", "esito ricerca",
    "fonte #", "titolo fonte", "URL fonte", "estratto fonte",
    "testo del sito da confrontare", "esito verifica contenuti", "stato sito ufficiale",
    "URL candidato sito ufficiale", "nota verifica sito", "azione"
]
research.append(research_headers)
for row in research_rows:
    research.append([row.get(header, "") for header in research_headers])

# Commercial prospect list: only businesses for which the search found no clear official website.
prospects = wb.create_sheet("Prospect senza sito")
prospects.append(["attività", "comune", "settore", "stato verifica", "query ricerca", "fonte emersa", "nota", "URL sito demo", "prossima azione"])
for slug, info in sorted(website_audit.items(), key=lambda kv: (kv[1]["comune"], kv[1]["name"])):
    if not info["status"].startswith("NESSUN SITO UFFICIALE"):
        continue
    source_url = info["url"]
    prospects.append([info["name"], info["comune"], info["settore"], info["status"], info["query"], source_url,
                      info["note"], "https://w0lw.github.io/varese-sites/sites/" + slug + "/",
                      "Controllo manuale Google/Maps e social; se confermato senza sito, contatto commerciale."])

with_site = wb.create_sheet("Da escludere - sito probabile")
with_site.append(["attività", "comune", "settore", "stato verifica", "URL candidato", "nota", "query ricerca", "URL sito demo"])
for slug, info in sorted(website_audit.items(), key=lambda kv: (kv[1]["comune"], kv[1]["name"])):
    if not info["status"].startswith("SITO WEB PROBABILE"):
        continue
    with_site.append([info["name"], info["comune"], info["settore"], info["status"], info["url"], info["note"],
                      info["query"], "https://w0lw.github.io/varese-sites/sites/" + slug + "/"])

corr_sheet = wb.create_sheet("Correzioni template")
if corrections:
    corr_sheet.append(list(corrections[0].keys()))
    for row in corrections:
        corr_sheet.append([row.get(key, "") for key in corrections[0].keys()])
else:
    corr_sheet.append(["Nessuna correzione template registrata"])

templates = wb.create_sheet("Distribuzione template")
templates.append(["Template ID", "Numero pagine", "Titolo/label"])
for key, count in sorted(template_counts.items(), key=lambda item: (-item[1], item[0])):
    label = next((clean(p.get("templateTitle")) for p in pages if clean(p.get("template")) == key and clean(p.get("templateTitle"))), "")
    templates.append([key, count, label])

for ws in wb.worksheets:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="20283A")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        max_len = 12
        for cell in ws[letter][:min(ws.max_row, 120)]:
            max_len = max(max_len, min(len(str(cell.value or "")), 42))
        ws.column_dimensions[letter].width = min(max_len + 2, 48)
    ws.sheet_view.showGridLines = False

for ws_name, widths in {
    "Riepilogo": {"A": 58, "B": 90},
    "Attività": {"A": 32, "B": 32, "C": 24, "D": 42, "E": 42, "F": 34, "G": 30, "H": 54, "I": 60, "J": 38, "K": 85, "L": 38, "M": 54, "N": 75, "O": 32, "P": 25, "Q": 75},
    "Ricerca web": {"A": 32, "B": 32, "C": 24, "D": 40, "E": 58, "F": 22, "G": 40, "H": 10, "I": 48, "J": 58, "K": 75, "L": 85, "M": 48, "N": 38, "O": 58, "P": 75, "Q": 75},
    "Prospect senza sito": {"A": 34, "B": 24, "C": 40, "D": 60, "E": 60, "F": 58, "G": 80, "H": 60, "I": 75},
    "Da escludere - sito probabile": {"A": 34, "B": 24, "C": 40, "D": 60, "E": 65, "F": 80, "G": 60, "H": 60},
    "Correzioni template": {"A": 42, "B": 42, "C": 28, "D": 34, "E": 34, "F": 24, "G": 80},
}.items():
    for col, width in widths.items():
        wb[ws_name].column_dimensions[col].width = width

wb.save(OUT)
print(json.dumps({
    "output": str(OUT),
    "pages": len(pages),
    "prospects_without_official_site_found": sum(1 for v in website_audit.values() if v["status"].startswith("NESSUN SITO UFFICIALE")),
    "probable_existing_websites": sum(1 for v in website_audit.values() if v["status"].startswith("SITO WEB PROBABILE")),
    "manual_review": sum(1 for v in website_audit.values() if v["status"].startswith(("DA VERIFICARE", "INCERTO"))),
    "searched_this_run": searched_this_run,
    "businesses_with_results": businesses_with_results,
    "businesses_with_no_results": businesses_with_no_results,
    "businesses_with_errors": businesses_with_errors,
    "sources": source_count,
    "search_backend": "Brave Search API" if API_KEY else ("DuckDuckGo (DDGS)" if DDGS is not None else "none"),
    "pages_with_flags": sum(1 for page in pages if clean(page.get("flags"))),
}, ensure_ascii=False, indent=2))
